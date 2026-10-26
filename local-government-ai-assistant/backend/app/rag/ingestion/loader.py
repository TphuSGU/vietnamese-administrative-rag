"""
Document loader for RAG LLM Core System.

Loads documents from PDF, DOCX, and TXT formats and extracts text content.
Implements Requirement 1 (Document Ingestion) from requirements document.
"""
import io
import shutil
from pathlib import Path
from datetime import datetime
from typing import Optional, Union
import logging

# Document processing libraries
import fitz  # PyMuPDF
import pdfplumber
from PIL import Image
from docx import Document as DocxDocument

try:
    import pytesseract
except ImportError:  # pragma: no cover - dependency is declared in requirements
    pytesseract = None

from app.core.config import settings
from app.models.data_models import Document
from app.models.exceptions import (
    DocumentLoadError,
    UnsupportedFormatError,
    ExtractionError
)
from app.core.logging import get_logger


logger = get_logger(__name__)


class DocumentLoader:
    """
    Loads and extracts text from documents in various formats.
    
    Implements:
    - Requirement 1.1: Accept PDF, DOCX, and TXT formats
    - Requirement 1.2: Extract text with UTF-8 encoding
    - Requirement 1.3: Error handling with descriptive messages
    - Requirement 1.4: Preserve metadata (filename, timestamp, file_size, format)
    - Requirement 15.4: Performance logging
    """
    
    SUPPORTED_FORMATS = {'.pdf', '.docx', '.txt'}
    
    def __init__(
        self,
        ocr_enabled: Optional[bool] = None,
        ocr_lang: Optional[str] = None,
        ocr_dpi: Optional[int] = None,
        ocr_min_chars_per_page: Optional[int] = None,
        ocr_max_pages: Optional[int] = None
    ):
        """
        Initialize document loader.

        OCR parameters default to application settings so callers (pipeline,
        tests) can override them without mutating global config.
        """
        self.logger = logger
        self.ocr_enabled = settings.ocr_enabled if ocr_enabled is None else ocr_enabled
        self.ocr_lang = settings.ocr_lang if ocr_lang is None else ocr_lang
        self.ocr_dpi = settings.ocr_dpi if ocr_dpi is None else ocr_dpi
        self.ocr_min_chars_per_page = (
            settings.ocr_min_chars_per_page
            if ocr_min_chars_per_page is None
            else ocr_min_chars_per_page
        )
        self.ocr_max_pages = settings.ocr_max_pages if ocr_max_pages is None else ocr_max_pages

        # Checked once per loader; tests can force this flag without the binary.
        self._tesseract_available = (
            pytesseract is not None and shutil.which("tesseract") is not None
        )
        if self.ocr_enabled and not self._tesseract_available:
            self.logger.warning(
                "OCR is enabled but tesseract/pytesseract is unavailable; "
                "scanned PDFs will fail with 'No text content extracted'",
                extra={"ocr_lang": self.ocr_lang}
            )
    
    def load_document(
        self,
        file_path: Optional[str] = None,
        file_content: Optional[bytes] = None,
        filename: Optional[str] = None,
        document_id: Optional[str] = None
    ) -> Document:
        """
        Load document from file path or file content.
        
        Implements Requirements 1.1, 1.2, 1.3, 1.4, 15.4.
        
        Args:
            file_path: Path to document file (for file-based loading)
            file_content: Binary content of file (for upload handling)
            filename: Original filename (required if using file_content)
            document_id: Optional document ID (generated if not provided)
            
        Returns:
            Document object with extracted text and metadata
            
        Raises:
            DocumentLoadError: If document cannot be loaded
            UnsupportedFormatError: If file format is not supported
            ExtractionError: If text extraction fails
        """
        start_time = datetime.utcnow()
        
        # Validate inputs
        if file_path is None and file_content is None:
            raise DocumentLoadError("Either file_path or file_content must be provided")
        
        if file_content is not None and filename is None:
            raise DocumentLoadError("filename is required when using file_content")
        
        # Determine filename and file extension
        if file_path:
            file_path_obj = Path(file_path)
            filename = filename or file_path_obj.name
            file_ext = file_path_obj.suffix.lower()
            
            # Read file content
            try:
                with open(file_path, 'rb') as f:
                    file_content = f.read()
            except Exception as e:
                self.logger.error(
                    f"Failed to read file: {file_path}",
                    extra={'error': str(e), 'file_name': filename}
                )
                raise DocumentLoadError(f"Failed to read file '{filename}': {str(e)}")
        else:
            file_ext = Path(filename).suffix.lower()
        
        # Validate file format (Requirement 1.1)
        if file_ext not in self.SUPPORTED_FORMATS:
            supported = ", ".join(sorted(self.SUPPORTED_FORMATS))
            raise UnsupportedFormatError(
                f"Unsupported file format '{file_ext}'. Supported formats: {supported}"
            )
        
        # Extract text based on format
        try:
            if file_ext == '.pdf':
                text = self._extract_pdf(file_content, filename)
            elif file_ext == '.docx':
                text = self._extract_docx(file_content, filename)
            elif file_ext == '.txt':
                text = self._extract_txt(file_content, filename)
            else:
                raise UnsupportedFormatError(f"Unsupported format: {file_ext}")
        except (UnsupportedFormatError, ExtractionError):
            raise
        except Exception as e:
            self.logger.error(
                f"Unexpected error extracting text from {filename}",
                extra={'error': str(e), 'format': file_ext}
            )
            raise ExtractionError(
                f"Failed to extract text from '{filename}': {str(e)}"
            )
        
        # Generate document ID if not provided
        if document_id is None:
            document_id = self._generate_document_id(filename)
        
        # Create Document object (Requirement 1.4)
        file_size = len(file_content)
        upload_timestamp = datetime.utcnow()
        
        document = Document(
            id=document_id,
            text=text,
            filename=filename,
            file_size=file_size,
            upload_timestamp=upload_timestamp,
            format=file_ext[1:]  # Remove leading dot
        )
        
        # Performance logging (Requirement 15.4)
        processing_time = (datetime.utcnow() - start_time).total_seconds()
        self.logger.info(
            f"Document loaded successfully: {filename}",
            extra={
                'document_id': document_id,
                'format': file_ext,
                'file_size_bytes': file_size,
                'text_length': len(text),
                'processing_time_seconds': processing_time
            }
        )
        
        return document
    
    def _extract_pdf(self, file_content: bytes, filename: str) -> str:
        """
        Extract text from PDF file using PyMuPDF and pdfplumber.
        
        Implements Requirement 1.2 (UTF-8 encoding).
        
        Args:
            file_content: Binary content of PDF file
            filename: Original filename for error messages
            
        Returns:
            Extracted text content
            
        Raises:
            ExtractionError: If text extraction fails
        """
        try:
            # Primary extraction with PyMuPDF (fitz), page by page so that
            # image-only pages can be routed to OCR without losing order.
            pdf_document = fitz.open(stream=file_content, filetype="pdf")
            page_texts = [""] * pdf_document.page_count

            for page_num in range(pdf_document.page_count):
                page_texts[page_num] = pdf_document[page_num].get_text()

            # If no text extracted, try pdfplumber for tables
            if sum(len(t.strip()) for t in page_texts) < 100:
                self.logger.info(
                    f"PyMuPDF extracted minimal text, trying pdfplumber for {filename}"
                )
                plumber_pages = self._extract_pdf_with_pdfplumber(file_content)
                if plumber_pages and len(plumber_pages) == len(page_texts):
                    page_texts = [
                        plumber or current
                        for plumber, current in zip(plumber_pages, page_texts)
                    ]
            # OCR fallback for pages that still have too little text.
            if self.ocr_enabled and self._tesseract_available:
                pages_to_ocr = [
                    index
                    for index, text in enumerate(page_texts)
                    if len(text.strip()) < self.ocr_min_chars_per_page
                ]
                if self.ocr_max_pages > 0:
                    pages_to_ocr = pages_to_ocr[: self.ocr_max_pages]

                for index in pages_to_ocr:
                    ocr_text = self._ocr_page(
                        pdf_document[index], index, filename
                    )
                    if not ocr_text:
                        continue
                    existing = page_texts[index].strip()
                    page_texts[index] = (
                        f"{existing}\n\n{ocr_text}" if existing else ocr_text
                    )

                if pages_to_ocr:
                    self.logger.info(
                        f"OCR fallback processed {len(pages_to_ocr)} page(s) "
                        f"for {filename}",
                        extra={
                            "ocr_lang": self.ocr_lang,
                            "ocr_dpi": self.ocr_dpi,
                            "ocr_pages": len(pages_to_ocr)
                        }
                    )

            pdf_document.close()

            text_parts = [text for text in page_texts if text.strip()]

        except Exception as e:
            self.logger.error(
                f"Failed to extract PDF content from {filename}",
                extra={'error': str(e)}
            )
            raise ExtractionError(
                f"Failed to extract text from PDF '{filename}': {str(e)}"
            )
        
        # Combine text parts
        full_text = "\n\n".join(text_parts)
        
        # Ensure UTF-8 encoding (Requirement 1.2)
        try:
            full_text.encode('utf-8')
        except UnicodeEncodeError:
            # Remove non-UTF-8 characters
            full_text = full_text.encode('utf-8', errors='ignore').decode('utf-8')
        
        if not full_text.strip():
            raise ExtractionError(
                f"No text content extracted from PDF '{filename}'"
            )
        
        return full_text
    
    def _extract_pdf_with_pdfplumber(self, file_content: bytes) -> list:
        """
        Extract text from PDF using pdfplumber (better for tables).
        
        Args:
            file_content: Binary content of PDF file
            
        Returns:
            List of text strings from each page
        """
        text_parts = []
        
        with pdfplumber.open(io.BytesIO(file_content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                # Keep one entry per page so OCR page indices stay aligned.
                text_parts.append(page_text or "")
        
        return text_parts

    def _ocr_page(self, page, page_num: int, filename: str) -> str:
        """
        Render a single PDF page and OCR it with Tesseract.

        Returns an empty string on failure so a partially scanned document can
        still be ingested from the pages that do have a text layer.
        """
        try:
            pixmap = page.get_pixmap(
                dpi=self.ocr_dpi,
                colorspace=fitz.csRGB,
                alpha=False
            )
            image = Image.frombytes(
                "RGB",
                (pixmap.width, pixmap.height),
                pixmap.samples
            )
            text = pytesseract.image_to_string(image, lang=self.ocr_lang)
            return text.strip()
        except Exception as exc:
            self.logger.warning(
                f"OCR failed for page {page_num + 1} of {filename}: {exc}. "
                f"Install tesseract-ocr and tesseract-ocr-{self.ocr_lang}.",
                extra={
                    "file_name": filename,
                    "page": page_num + 1,
                    "ocr_lang": self.ocr_lang,
                    "error": str(exc)
                }
            )
            return ""

    def _extract_docx(self, file_content: bytes, filename: str) -> str:
        """
        Extract text from DOCX file.
        
        Implements Requirement 1.2 (UTF-8 encoding).
        
        Args:
            file_content: Binary content of DOCX file
            filename: Original filename for error messages
            
        Returns:
            Extracted text content
            
        Raises:
            ExtractionError: If text extraction fails
        """
        try:
            doc = DocxDocument(io.BytesIO(file_content))
            
            # Extract paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            
            full_text = "\n\n".join(paragraphs)
            
        except Exception as e:
            self.logger.error(
                f"Failed to extract DOCX content from {filename}",
                extra={'error': str(e)}
            )
            raise ExtractionError(
                f"Failed to extract text from DOCX '{filename}': {str(e)}"
            )
        
        # Ensure UTF-8 encoding (Requirement 1.2)
        try:
            full_text.encode('utf-8')
        except UnicodeEncodeError:
            full_text = full_text.encode('utf-8', errors='ignore').decode('utf-8')
        
        if not full_text.strip():
            raise ExtractionError(
                f"No text content extracted from DOCX '{filename}'"
            )
        
        return full_text
    
    def _extract_txt(self, file_content: bytes, filename: str) -> str:
        """
        Extract text from TXT file with UTF-8 encoding.
        
        Implements Requirement 1.2 (UTF-8 encoding).
        
        Args:
            file_content: Binary content of TXT file
            filename: Original filename for error messages
            
        Returns:
            Extracted text content
            
        Raises:
            ExtractionError: If text extraction fails
        """
        try:
            # Try UTF-8 first (Requirement 1.2)
            text = file_content.decode('utf-8')
        except UnicodeDecodeError:
            # Try other common encodings
            for encoding in ['utf-8-sig', 'latin-1', 'cp1252']:
                try:
                    text = file_content.decode(encoding)
                    self.logger.warning(
                        f"File {filename} decoded with {encoding} instead of UTF-8"
                    )
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ExtractionError(
                    f"Failed to decode TXT file '{filename}' with UTF-8 encoding"
                )
        
        if not text.strip():
            raise ExtractionError(
                f"No text content in TXT file '{filename}'"
            )
        
        return text
    
    def _generate_document_id(self, filename: str) -> str:
        """
        Generate unique document ID from filename and timestamp.
        
        Args:
            filename: Original filename
            
        Returns:
            Unique document ID
        """
        import hashlib
        
        timestamp = datetime.utcnow().isoformat()
        content = f"{filename}_{timestamp}"
        doc_id = hashlib.sha256(content.encode()).hexdigest()[:16]
        
        return doc_id
