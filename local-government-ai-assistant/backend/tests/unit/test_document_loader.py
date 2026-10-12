"""
Unit tests for DocumentLoader.

Tests Requirement 1 (Document Ingestion) from requirements document.
"""
import pytest
import io
from datetime import datetime
from pathlib import Path

from app.rag.ingestion.loader import DocumentLoader
from app.models.data_models import Document
from app.models.exceptions import (
    DocumentLoadError,
    UnsupportedFormatError,
    ExtractionError
)


@pytest.fixture
def loader():
    """Create DocumentLoader instance."""
    return DocumentLoader()


@pytest.fixture
def sample_txt_content():
    """Sample TXT file content."""
    return b"This is a test document.\nIt has multiple lines.\nVietnamese: Xin chao!"


@pytest.fixture
def sample_pdf_path(tmp_path):
    """Create a simple PDF file for testing."""
    import fitz
    
    pdf_path = tmp_path / "test.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((100, 100), "Test PDF Content\nLine 2\nVietnamese: Thủ tục hành chính")
    doc.save(pdf_path)
    doc.close()
    
    return str(pdf_path)


@pytest.fixture
def sample_docx_path(tmp_path):
    """Create a simple DOCX file for testing."""
    from docx import Document as DocxDocument
    
    docx_path = tmp_path / "test.docx"
    doc = DocxDocument()
    doc.add_paragraph("Test DOCX Content")
    doc.add_paragraph("Line 2")
    doc.add_paragraph("Vietnamese: Giấy phép kinh doanh")
    doc.save(docx_path)
    
    return str(docx_path)


class TestDocumentLoaderBasic:
    """Test basic document loading functionality."""
    
    def test_loader_initialization(self, loader):
        """DocumentLoader should initialize successfully."""
        assert loader is not None
        assert loader.SUPPORTED_FORMATS == {'.pdf', '.docx', '.txt'}
    
    def test_load_requires_input(self, loader):
        """Loading without file_path or file_content should raise error."""
        with pytest.raises(DocumentLoadError, match="file_path or file_content must be provided"):
            loader.load_document()
    
    def test_load_content_requires_filename(self, loader, sample_txt_content):
        """Loading with file_content but no filename should raise error."""
        with pytest.raises(DocumentLoadError, match="filename is required"):
            loader.load_document(file_content=sample_txt_content)


class TestTXTLoading:
    """Test TXT file loading (Requirement 1.1, 1.2)."""
    
    def test_load_txt_from_content(self, loader, sample_txt_content):
        """Should load TXT file from binary content."""
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt"
        )
        
        assert isinstance(doc, Document)
        assert "test document" in doc.text.lower()
        assert "multiple lines" in doc.text.lower()
        assert doc.filename == "test.txt"
        assert doc.format == "txt"
        assert doc.file_size == len(sample_txt_content)
    
    def test_load_txt_from_path(self, loader, tmp_path, sample_txt_content):
        """Should load TXT file from file path."""
        txt_path = tmp_path / "test.txt"
        txt_path.write_bytes(sample_txt_content)
        
        doc = loader.load_document(file_path=str(txt_path))
        
        assert isinstance(doc, Document)
        assert "test document" in doc.text.lower()
        assert doc.filename == "test.txt"
        assert doc.format == "txt"
    
    def test_txt_utf8_encoding(self, loader, tmp_path):
        """TXT files should be decoded with UTF-8 encoding (Requirement 1.2)."""
        vietnamese_content = "Thủ tục hành chính\nGiấy phép kinh doanh\nĐăng ký tạm trú"
        txt_path = tmp_path / "vietnamese.txt"
        txt_path.write_text(vietnamese_content, encoding='utf-8')
        
        doc = loader.load_document(file_path=str(txt_path))
        
        assert "Thủ tục" in doc.text
        assert "Giấy phép" in doc.text
        assert "Đăng ký" in doc.text
    
    def test_txt_empty_file(self, loader, tmp_path):
        """Empty TXT file should raise ExtractionError."""
        txt_path = tmp_path / "empty.txt"
        txt_path.write_text("")
        
        with pytest.raises(ExtractionError, match="No text content"):
            loader.load_document(file_path=str(txt_path))
    
    def test_txt_whitespace_only(self, loader, tmp_path):
        """Whitespace-only TXT file should raise ExtractionError."""
        txt_path = tmp_path / "whitespace.txt"
        txt_path.write_text("   \n\n\t   ")
        
        with pytest.raises(ExtractionError, match="No text content"):
            loader.load_document(file_path=str(txt_path))


class TestPDFLoading:
    """Test PDF file loading (Requirement 1.1, 1.2)."""
    
    def test_load_pdf_from_path(self, loader, sample_pdf_path):
        """Should load PDF file from file path."""
        doc = loader.load_document(file_path=sample_pdf_path)
        
        assert isinstance(doc, Document)
        assert "Test PDF Content" in doc.text
        assert doc.filename == "test.pdf"
        assert doc.format == "pdf"
        assert doc.file_size > 0
    
    def test_load_pdf_from_content(self, loader, sample_pdf_path):
        """Should load PDF file from binary content."""
        with open(sample_pdf_path, 'rb') as f:
            pdf_content = f.read()
        
        doc = loader.load_document(
            file_content=pdf_content,
            filename="document.pdf"
        )
        
        assert isinstance(doc, Document)
        assert "Test PDF Content" in doc.text
        assert doc.filename == "document.pdf"
        assert doc.format == "pdf"
    
    def test_pdf_multipage(self, loader, tmp_path):
        """Should extract text from multi-page PDF."""
        import fitz
        
        pdf_path = tmp_path / "multipage.pdf"
        doc = fitz.open()
        
        # Page 1
        page1 = doc.new_page()
        page1.insert_text((100, 100), "Page 1 content")
        
        # Page 2
        page2 = doc.new_page()
        page2.insert_text((100, 100), "Page 2 content")
        
        doc.save(pdf_path)
        doc.close()
        
        loaded_doc = loader.load_document(file_path=str(pdf_path))
        
        assert "Page 1 content" in loaded_doc.text
        assert "Page 2 content" in loaded_doc.text
    
    def test_pdf_vietnamese_text(self, loader, tmp_path):
        """Should handle Vietnamese text in PDF (Requirement 1.2)."""
        import fitz
        
        pdf_path = tmp_path / "vietnamese.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((100, 100), "Thủ tục đăng ký\nGiấy phép xây dựng")
        doc.save(pdf_path)
        doc.close()
        
        loaded_doc = loader.load_document(file_path=str(pdf_path))
        
        # Should contain Vietnamese text
        assert "Thủ tục" in loaded_doc.text or "thu tuc" in loaded_doc.text.lower()


class TestDOCXLoading:
    """Test DOCX file loading (Requirement 1.1, 1.2)."""
    
    def test_load_docx_from_path(self, loader, sample_docx_path):
        """Should load DOCX file from file path."""
        doc = loader.load_document(file_path=sample_docx_path)
        
        assert isinstance(doc, Document)
        assert "Test DOCX Content" in doc.text
        assert doc.filename == "test.docx"
        assert doc.format == "docx"
        assert doc.file_size > 0
    
    def test_load_docx_from_content(self, loader, sample_docx_path):
        """Should load DOCX file from binary content."""
        with open(sample_docx_path, 'rb') as f:
            docx_content = f.read()
        
        doc = loader.load_document(
            file_content=docx_content,
            filename="document.docx"
        )
        
        assert isinstance(doc, Document)
        assert "Test DOCX Content" in doc.text
        assert doc.filename == "document.docx"
        assert doc.format == "docx"
    
    def test_docx_multiple_paragraphs(self, loader, tmp_path):
        """Should extract multiple paragraphs from DOCX."""
        from docx import Document as DocxDocument
        
        docx_path = tmp_path / "multi_para.docx"
        doc = DocxDocument()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("Paragraph 2")
        doc.add_paragraph("Paragraph 3")
        doc.save(docx_path)
        
        loaded_doc = loader.load_document(file_path=str(docx_path))
        
        assert "Paragraph 1" in loaded_doc.text
        assert "Paragraph 2" in loaded_doc.text
        assert "Paragraph 3" in loaded_doc.text
    
    def test_docx_with_tables(self, loader, tmp_path):
        """Should extract text from tables in DOCX."""
        from docx import Document as DocxDocument
        
        docx_path = tmp_path / "with_table.docx"
        doc = DocxDocument()
        doc.add_paragraph("Before table")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cell 1"
        table.cell(0, 1).text = "Cell 2"
        table.cell(1, 0).text = "Cell 3"
        table.cell(1, 1).text = "Cell 4"
        
        doc.add_paragraph("After table")
        doc.save(docx_path)
        
        loaded_doc = loader.load_document(file_path=str(docx_path))
        
        assert "Before table" in loaded_doc.text
        assert "Cell 1" in loaded_doc.text
        assert "Cell 4" in loaded_doc.text
        assert "After table" in loaded_doc.text
    
    def test_docx_vietnamese_text(self, loader, tmp_path):
        """Should handle Vietnamese text in DOCX (Requirement 1.2)."""
        from docx import Document as DocxDocument
        
        docx_path = tmp_path / "vietnamese.docx"
        doc = DocxDocument()
        doc.add_paragraph("Thủ tục hành chính")
        doc.add_paragraph("Giấy phép kinh doanh")
        doc.save(docx_path)
        
        loaded_doc = loader.load_document(file_path=str(docx_path))
        
        assert "Thủ tục" in loaded_doc.text
        assert "Giấy phép" in loaded_doc.text
    
    def test_docx_empty(self, loader, tmp_path):
        """Empty DOCX should raise ExtractionError."""
        from docx import Document as DocxDocument
        
        docx_path = tmp_path / "empty.docx"
        doc = DocxDocument()
        doc.save(docx_path)
        
        with pytest.raises(ExtractionError, match="No text content"):
            loader.load_document(file_path=str(docx_path))


class TestUnsupportedFormats:
    """Test handling of unsupported file formats (Requirement 1.1)."""
    
    def test_unsupported_extension(self, loader, tmp_path):
        """Unsupported file extension should raise UnsupportedFormatError."""
        exe_path = tmp_path / "file.exe"
        exe_path.write_bytes(b"fake executable")
        
        with pytest.raises(UnsupportedFormatError, match="Unsupported file format"):
            loader.load_document(file_path=str(exe_path))
    
    def test_unsupported_extension_from_content(self, loader):
        """Unsupported extension in filename should raise error."""
        with pytest.raises(UnsupportedFormatError, match="Unsupported file format"):
            loader.load_document(
                file_content=b"content",
                filename="file.xyz"
            )


class TestDocumentMetadata:
    """Test document metadata preservation (Requirement 1.4)."""
    
    def test_metadata_filename(self, loader, sample_txt_content):
        """Document should preserve filename."""
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="my_document.txt"
        )
        
        assert doc.filename == "my_document.txt"
    
    def test_metadata_file_size(self, loader, sample_txt_content):
        """Document should preserve file size."""
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt"
        )
        
        assert doc.file_size == len(sample_txt_content)
    
    def test_metadata_format(self, loader, sample_txt_content):
        """Document should preserve file format."""
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt"
        )
        
        assert doc.format == "txt"
    
    def test_metadata_timestamp(self, loader, sample_txt_content):
        """Document should have upload timestamp."""
        before = datetime.utcnow()
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt"
        )
        after = datetime.utcnow()
        
        assert before <= doc.upload_timestamp <= after
    
    def test_metadata_document_id(self, loader, sample_txt_content):
        """Document should have unique ID."""
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt"
        )
        
        assert doc.id is not None
        assert len(doc.id) > 0
    
    def test_custom_document_id(self, loader, sample_txt_content):
        """Should accept custom document ID."""
        custom_id = "custom_doc_123"
        doc = loader.load_document(
            file_content=sample_txt_content,
            filename="test.txt",
            document_id=custom_id
        )
        
        assert doc.id == custom_id


class TestErrorHandling:
    """Test error handling (Requirement 1.3)."""
    
    def test_file_not_found(self, loader):
        """Non-existent file should raise DocumentLoadError."""
        with pytest.raises(DocumentLoadError, match="Failed to read file"):
            loader.load_document(file_path="/nonexistent/file.pdf")
    
    def test_corrupted_pdf(self, loader, tmp_path):
        """Corrupted PDF should raise ExtractionError."""
        pdf_path = tmp_path / "corrupted.pdf"
        pdf_path.write_bytes(b"This is not a valid PDF")
        
        with pytest.raises(ExtractionError, match="Failed to extract text"):
            loader.load_document(file_path=str(pdf_path))
    
    def test_corrupted_docx(self, loader, tmp_path):
        """Corrupted DOCX should raise ExtractionError."""
        docx_path = tmp_path / "corrupted.docx"
        docx_path.write_bytes(b"This is not a valid DOCX")
        
        with pytest.raises(ExtractionError, match="Failed to extract text"):
            loader.load_document(file_path=str(docx_path))
    
    def test_error_message_includes_filename(self, loader, tmp_path):
        """Error messages should include filename (Requirement 1.3)."""
        pdf_path = tmp_path / "bad_file.pdf"
        pdf_path.write_bytes(b"invalid")
        
        try:
            loader.load_document(file_path=str(pdf_path))
            assert False, "Should have raised ExtractionError"
        except ExtractionError as e:
            assert "bad_file.pdf" in str(e)
